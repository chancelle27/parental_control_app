import json
import os
import sys
import csv
import numpy as np
from PyQt5.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton, QTableWidget,
    QTableWidgetItem, QHeaderView, QComboBox, QDateEdit, QColorDialog,
    QApplication, QFileDialog, QMessageBox
)
from PyQt5.QtGui import QFont, QColor
from PyQt5.QtCore import Qt, QPropertyAnimation, QEasingCurve, QDate
import pyqtgraph as pg
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

class ReportPage(QWidget):
    def __init__(self):
        super().__init__()
        # Charger les données depuis le fichier JSON
        self.usage_data = self.load_usage_data("parental_control_app/data/usage_data.json")
        # Valeur par défaut pour la couleur des graphiques
        self.graph_color = '#3498db'
        self.init_ui()
        self.setup_animations()
        self.apply_filters()

    def load_usage_data(self, file_path):
        # Utilise sys.argv[0] pour obtenir le répertoire d'exécution
        base_path = os.path.dirname(os.path.abspath(sys.argv[0]))
        # Construire le chemin vers le dossier 'data' qui se trouve dans le dossier racine de l'application
        full_path = os.path.join(base_path, "data", "usage_data.json")
        with open(full_path, "r") as f:
            return json.load(f)

    def init_ui(self):
        self.setWindowTitle("Rapport d'Activité - Filtrage Dynamique")
        layout = QVBoxLayout(self)
        layout.setSpacing(20)
        layout.setContentsMargins(30, 30, 30, 30)

        # Titre du rapport
        title = QLabel("Rapport d'Activité")
        title.setFont(QFont("Arial", 24, QFont.Bold))
        title.setStyleSheet("color: #2c3e50; margin-bottom: 20px;")
        title.setAlignment(Qt.AlignCenter)
        layout.addWidget(title)

        # Section des filtres
        filter_layout = QHBoxLayout()
        filter_layout.setSpacing(20)

        # Filtre par application
        self.app_filter = QComboBox()
        self.app_filter.addItem("Toutes les applications")
        for app in sorted(self.usage_data.keys()):
            self.app_filter.addItem(app)
        self.app_filter.currentIndexChanged.connect(self.apply_filters)
        filter_layout.addWidget(QLabel("Application:"))
        filter_layout.addWidget(self.app_filter)

        # Filtre par date
        self.date_filter = QDateEdit()
        self.date_filter.setCalendarPopup(True)
        self.date_filter.setDate(QDate.currentDate())
        self.date_filter.dateChanged.connect(self.apply_filters)
        filter_layout.addWidget(QLabel("Date:"))
        filter_layout.addWidget(self.date_filter)

        # Bouton pour choisir la couleur du graphique
        self.color_button = QPushButton("Choisir Couleur Graphique")
        self.color_button.clicked.connect(self.choose_color)
        filter_layout.addWidget(self.color_button)

        layout.addLayout(filter_layout)

        # Section des graphiques
        charts_layout = QHBoxLayout()
        charts_layout.setSpacing(20)
        
        # Graphique linéaire (avec style amélioré)
        self.line_chart = pg.PlotWidget(title="Graphique Dynamique")
        self.line_chart.setBackground('#ffffff')  # Fond blanc pour un rendu épuré
        self.line_chart.showGrid(x=True, y=True, alpha=0.3)
        self.line_chart.getAxis('left').setPen(pg.mkPen(color="#2c3e50", width=2))
        self.line_chart.getAxis('bottom').setPen(pg.mkPen(color="#2c3e50", width=2))
        self.line_chart.getAxis('left').setTextPen(pg.mkPen(color="#2c3e50"))
        self.line_chart.getAxis('bottom').setTextPen(pg.mkPen(color="#2c3e50"))
        self.line_chart.setLabel('left', 'Temps (min)', **{'color': '#2c3e50', 'font-size': '14pt'})
        self.line_chart.setLabel('bottom', 'Indice / Date', **{'color': '#2c3e50', 'font-size': '14pt'})
        charts_layout.addWidget(self.line_chart)
        
        # Graphique à barres (avec style amélioré)
        self.bar_chart = pg.PlotWidget(title="Détails par Application")
        self.bar_chart.setBackground('#ffffff')
        self.bar_chart.showGrid(x=True, y=True, alpha=0.3)
        self.bar_chart.getAxis('left').setPen(pg.mkPen(color="#2c3e50", width=2))
        self.bar_chart.getAxis('bottom').setPen(pg.mkPen(color="#2c3e50", width=2))
        self.bar_chart.getAxis('left').setTextPen(pg.mkPen(color="#2c3e50"))
        self.bar_chart.getAxis('bottom').setTextPen(pg.mkPen(color="#2c3e50"))
        self.bar_chart.setLabel('left', 'Temps (min)', **{'color': '#2c3e50', 'font-size': '14pt'})
        self.bar_chart.setLabel('bottom', 'Applications', **{'color': '#2c3e50', 'font-size': '14pt'})
        charts_layout.addWidget(self.bar_chart)
        
        layout.addLayout(charts_layout)

        # Tableau récapitulatif
        summary_title = QLabel("Détails de l'Activité")
        summary_title.setFont(QFont("Arial", 18, QFont.Bold))
        summary_title.setStyleSheet("color: #2c3e50;")
        summary_title.setAlignment(Qt.AlignCenter)
        layout.addWidget(summary_title)

        self.summary_table = QTableWidget(0, 3)
        self.summary_table.setHorizontalHeaderLabels(["Application", "Temps Utilisé (min)", "Limite (min)"])
        self.summary_table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        layout.addWidget(self.summary_table)

        # Bouton pour exporter le rapport
        export_button = QPushButton("Exporter le Rapport")
        export_button.setFont(QFont("Arial", 14))
        export_button.setStyleSheet("""
            QPushButton {
                background-color: #1abc9c;
                color: white;
                padding: 10px 20px;
                border-radius: 5px;
            }
            QPushButton:hover {
                background-color: #16a085;
            }
        """)
        export_button.clicked.connect(self.export_report)
        layout.addWidget(export_button, alignment=Qt.AlignCenter)

    def setup_animations(self):
        # Animation de fondu pour le tableau
        animation = QPropertyAnimation(self.summary_table, b"windowOpacity")
        animation.setDuration(1000)
        animation.setStartValue(0)
        animation.setEndValue(1)
        animation.setEasingCurve(QEasingCurve.InOutQuad)
        animation.start()

    def choose_color(self):
        """Ouvre une boîte de dialogue pour choisir la couleur du graphique."""
        color = QColorDialog.getColor()
        if color.isValid():
            self.graph_color = color.name()
            self.apply_filters()

    def apply_filters(self):
        """Applique les filtres et met à jour les graphiques et le tableau."""
        selected_app = self.app_filter.currentText()
        selected_date = self.date_filter.date().toString("yyyy-MM-dd")

        # Mise à jour du graphique linéaire :
        self.line_chart.clear()
        if selected_app == "Toutes les applications":
            daily_totals = {}
            for app, data in self.usage_data.items():
                for date, usage in data["days"].items():
                    daily_totals[date] = daily_totals.get(date, 0) + usage
            dates = sorted(daily_totals.keys())
            x_vals = np.arange(len(dates))
            y_vals = np.array([daily_totals[date] for date in dates])
            pen = pg.mkPen(color=self.graph_color, width=2)
            self.line_chart.plot(x_vals, y_vals, pen=pen, symbol='o', symbolSize=8, symbolBrush=self.graph_color)
            ticks = list(zip(x_vals, dates))
            self.line_chart.getAxis('bottom').setTicks([ticks])
            self.line_chart.setTitle("Évolution Globale du Temps d'Écran", color="#2c3e50", size="16pt")
        else:
            app_data = self.usage_data.get(selected_app, {})
            dates = sorted(app_data.get("days", {}).keys())
            x_vals = np.arange(len(dates))
            y_vals = np.array([app_data["days"][date] for date in dates])
            pen = pg.mkPen(color=self.graph_color, width=2)
            self.line_chart.plot(x_vals, y_vals, pen=pen, symbol='o', symbolSize=8, symbolBrush=self.graph_color)
            ticks = list(zip(x_vals, dates))
            self.line_chart.getAxis('bottom').setTicks([ticks])
            self.line_chart.setTitle(f"Évolution de {selected_app}", color="#2c3e50", size="16pt")

        # Mise à jour du graphique à barres :
        self.bar_chart.clear()
        if selected_app == "Toutes les applications":
            apps = []
            usage_values = []
            for app, data in self.usage_data.items():
                apps.append(app)
                usage = data.get("days", {}).get(selected_date, 0)
                usage_values.append(usage)
            indices = np.arange(len(apps))
            bg = pg.BarGraphItem(x=indices, height=usage_values, width=0.6, brush=self.graph_color)
            self.bar_chart.addItem(bg)
            ticks = list(zip(indices, apps))
            self.bar_chart.getAxis('bottom').setTicks([ticks])
            self.bar_chart.setTitle(f"Utilisation par Application le {selected_date}", color="#2c3e50", size="16pt")
        else:
            data = self.usage_data.get(selected_app, {})
            usage = data.get("days", {}).get(selected_date, 0)
            limit = data.get("limit", 0)
            apps = [selected_app, selected_app]
            values = [usage, limit]
            indices = np.arange(len(apps))
            bg = pg.BarGraphItem(x=indices, height=values, width=0.6, brush=self.graph_color)
            self.bar_chart.addItem(bg)
            ticks = list(zip(indices, ["Utilisé", "Limite"]))
            self.bar_chart.getAxis('bottom').setTicks([ticks])
            self.bar_chart.setTitle(f"{selected_app} le {selected_date}", color="#2c3e50", size="16pt")

        # Mise à jour du tableau de synthèse
        self.update_table(selected_app, selected_date)

    def update_table(self, selected_app, selected_date):
        """Met à jour le tableau avec les données filtrées."""
        if selected_app == "Toutes les applications":
            filtered_data = []
            for app, data in self.usage_data.items():
                usage = data.get("days", {}).get(selected_date, 0)
                limit = data.get("limit", 0)
                filtered_data.append((app, usage, limit))
        else:
            data = self.usage_data.get(selected_app, {})
            usage = data.get("days", {}).get(selected_date, 0)
            limit = data.get("limit", 0)
            filtered_data = [(selected_app, usage, limit)]
        self.summary_table.setRowCount(len(filtered_data))
        for row, (app, usage, limit) in enumerate(filtered_data):
            self.summary_table.setItem(row, 0, QTableWidgetItem(app))
            self.summary_table.setItem(row, 1, QTableWidgetItem(str(usage)))
            self.summary_table.setItem(row, 2, QTableWidgetItem(str(limit)))

    def export_report(self):
        """Exporte les données affichées dans le tableau récapitulatif au format Excel (XLSX)."""
        # Ouvrir une boîte de dialogue pour sélectionner le fichier de destination
        file_path, _ = QFileDialog.getSaveFileName(
            self, "Exporter le Rapport", "", "Excel Files (*.xlsx);;All Files (*)"
        )
        if not file_path:
            return  # L'utilisateur a annulé

        try:
            # Créer un nouveau classeur Excel
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Rapport d'Activité"
            
            # Définir un style pour les en-têtes
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill("solid", fgColor="1abc9c")
            header_alignment = Alignment(horizontal="center", vertical="center")
            thin_border = Border(left=Side(style="thin"), 
                                 right=Side(style="thin"),
                                 top=Side(style="thin"),
                                 bottom=Side(style="thin"))
            
            # Récupérer les en-têtes du tableau
            headers = [self.summary_table.horizontalHeaderItem(col).text() for col in range(self.summary_table.columnCount())]
            ws.append(headers)
            for col in range(1, len(headers)+1):
                cell = ws.cell(row=1, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = thin_border

            # Récupérer les données du tableau et les écrire
            for row in range(self.summary_table.rowCount()):
                row_data = []
                for col in range(self.summary_table.columnCount()):
                    item = self.summary_table.item(row, col)
                    row_data.append(item.text() if item is not None else "")
                ws.append(row_data)
            
            # Ajuster la largeur des colonnes
            for col in ws.columns:
                max_length = 0
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                for cell in col:
                    try:
                        if cell.value:
                            max_length = max(max_length, len(str(cell.value)))
                    except Exception:
                        pass
                adjusted_width = (max_length + 2)
                ws.column_dimensions[col_letter].width = adjusted_width

            # Sauvegarder le fichier Excel
            wb.save(file_path)
            QMessageBox.information(self, "Export Réussi", f"Le rapport a été exporté avec succès vers :\n{file_path}")
        except Exception as e:
            QMessageBox.critical(self, "Erreur d'Export", f"Une erreur est survenue lors de l'export du rapport :\n{str(e)}")


if __name__ == "__main__":
    app = QApplication(sys.argv)
    report = ReportPage()
    report.resize(1200, 800)
    report.show()
    sys.exit(app.exec_())
